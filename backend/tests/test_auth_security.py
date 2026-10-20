from datetime import datetime, timedelta, timezone

import jwt

from app.core.config import get_settings
from app.core.security import create_access_token, hash_password
from app.models.listing import Listing
from app.models.location import Location
from app.models.property import Property, PropertyType
from app.models.user import User


def _user(db, email, role="agent", password="Secure-User-Passphrase-42!", active=True):
    user = User(email=email, username=email.split("@")[0], full_name=email.split("@")[0], role=role,
                password_hash=hash_password(password), is_active=active)
    db.add(user)
    db.commit()
    db.refresh(user)
    return user


def test_protected_api_rejects_missing_invalid_and_expired_tokens(anonymous_client, db_session):
    assert anonymous_client.get("/api/v1/listings").status_code == 401
    assert anonymous_client.get("/api/v1/listings", headers={"Authorization": "Bearer not-a-jwt"}).status_code == 401
    settings = get_settings()
    expired = jwt.encode({
        "sub": "1", "iat": datetime.now(timezone.utc) - timedelta(hours=2),
        "exp": datetime.now(timezone.utc) - timedelta(hours=1),
    }, settings.jwt_secret_key, algorithm="HS256")
    assert anonymous_client.get("/api/v1/listings", headers={"Authorization": f"Bearer {expired}"}).status_code == 401
    wrong_signature = jwt.encode({
        "sub": "1", "iat": datetime.now(timezone.utc),
        "exp": datetime.now(timezone.utc) + timedelta(minutes=5),
    }, "wrong-signature-key-that-is-long-enough-for-hs256", algorithm="HS256")
    assert anonymous_client.get("/api/v1/listings", headers={"Authorization": f"Bearer {wrong_signature}"}).status_code == 401


def test_login_returns_signed_access_token_and_me(anonymous_client, db_session):
    user = _user(db_session, "login-agent@example.com")
    response = anonymous_client.post("/api/v1/auth/login", json={
        "username": user.username, "password": "Secure-User-Passphrase-42!"
    })
    assert response.status_code == 200, response.text
    token = response.json()["access_token"]
    assert response.json()["user"]["role"] == "agent"
    assert "rea_access_token" in response.cookies
    me = anonymous_client.get("/api/v1/auth/me", headers={"Authorization": f"Bearer {token}"})
    assert me.status_code == 200
    assert me.json()["email"] == user.email


def test_login_does_not_disclose_unknown_or_wrong_password(anonymous_client, db_session):
    _user(db_session, "known-agent@example.com")
    wrong = anonymous_client.post("/api/v1/auth/login", json={
        "username": "known-agent", "password": "incorrect-password"
    })
    unknown = anonymous_client.post("/api/v1/auth/login", json={
        "username": "unknown-agent", "password": "incorrect-password"
    })
    assert wrong.status_code == unknown.status_code == 401
    assert wrong.json()["detail"] == unknown.json()["detail"]


def test_password_hash_is_not_an_accepted_user_input(client):
    response = client.post("/api/v1/users", json={
        "email": "bad-hash@example.com", "password": "Secure-User-Passphrase-42!",
        "password_hash": "attacker-supplied-hash", "full_name": "Bad Hash", "role": "agent",
    })
    assert response.status_code == 422
    assert "attacker-supplied-hash" not in response.text


def test_agent_cannot_access_admin_user_management(anonymous_client, db_session):
    user = _user(db_session, "ordinary-agent@example.com", role="agent")
    login = anonymous_client.post("/api/v1/auth/login", json={
        "username": user.username, "password": "Secure-User-Passphrase-42!"
    })
    token = login.json()["access_token"]
    response = anonymous_client.get("/api/v1/users", headers={"Authorization": f"Bearer {token}"})
    assert response.status_code == 403


def test_agent_cannot_update_or_archive_another_agents_listing(anonymous_client, db_session):
    location = Location(level="neighborhood", name="Moda", normalized_name="moda")
    ptype = PropertyType(id=1, code="apartment", name="Daire", group_code="residential")
    owner = _user(db_session, "listing-owner@example.com", role="agent")
    actor = _user(db_session, "listing-actor@example.com", role="agent")
    db_session.add_all([location, ptype])
    db_session.commit()
    prop = Property(property_type_id=ptype.id, location_id=location.id)
    db_session.add(prop)
    db_session.commit()
    listing = Listing(property_id=prop.id, created_by=owner.id, asking_price=1000000, currency="TRY")
    db_session.add(listing)
    db_session.commit()
    db_session.refresh(listing)
    login = anonymous_client.post("/api/v1/auth/login", json={
        "username": actor.username, "password": "Secure-User-Passphrase-42!"
    })
    token = login.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    updated = anonymous_client.patch(f"/api/v1/listings/{listing.id}", json={"asking_price": "900000"}, headers=headers)
    archived = anonymous_client.delete(f"/api/v1/listings/{listing.id}", headers=headers)
    assert updated.status_code == 403
    assert archived.status_code == 403
    db_session.refresh(listing)
    assert str(listing.asking_price) == "1000000.00"
    assert listing.archived_at is None


def test_agent_can_create_listing_but_server_sets_owner(client, db_session):
    location = Location(level="neighborhood", name="Test", normalized_name="test")
    ptype = PropertyType(id=1, code="apartment", name="Daire", group_code="residential")
    other_user = User(email="other@example.com", username="other", password_hash=hash_password("Other-User-Passphrase-42!"), full_name="Other", role="agent")
    db_session.add_all([location, ptype, other_user])
    db_session.commit()
    prop = Property(property_type_id=ptype.id, location_id=location.id)
    db_session.add(prop)
    db_session.commit()
    response = client.post("/api/v1/listings", json={"property_id": prop.id, "created_by": other_user.id, "asking_price": "1230000"})
    assert response.status_code == 201, response.text
    assert response.json()["created_by"] != other_user.id


def test_agent_cannot_reimport_another_agents_listing(anonymous_client, db_session):
    location = Location(level="neighborhood", name="Import Moda", normalized_name="import-moda")
    ptype = PropertyType(id=1, code="apartment", name="Daire", group_code="residential")
    owner = _user(db_session, "import-owner@example.com", role="agent")
    actor = _user(db_session, "import-actor@example.com", role="agent")
    db_session.add_all([location, ptype])
    db_session.commit()
    prop = Property(property_type_id=ptype.id, location_id=location.id)
    db_session.add(prop)
    db_session.commit()
    listing = Listing(property_id=prop.id, created_by=owner.id, source_type="authorized_source",
                      source_name="example.org", source_listing_id="listing-555",
                      source_url="https://example.org/ilan/555", asking_price=1000000, currency="TRY")
    db_session.add(listing)
    db_session.commit()
    login = anonymous_client.post("/api/v1/auth/login", json={
        "username": actor.username, "password": "Secure-User-Passphrase-42!"
    })
    token = login.json()["access_token"]
    response = anonymous_client.post("/api/v1/imports/listing", headers={"Authorization": f"Bearer {token}"}, json={
        "source_url": "https://example.org/ilan/555", "source_name": "example.org",
        "source_listing_id": "listing-555", "location_id": location.id,
        "property_type_code": "apartment", "asking_price": "900000"
    })
    assert response.status_code == 403
    db_session.refresh(listing)
    assert str(listing.asking_price) == "1000000.00"


def test_logout_revokes_bearer_token_server_side(anonymous_client, db_session):
    user = _user(db_session, "logout-agent@example.com")
    login = anonymous_client.post("/api/v1/auth/login", json={
        "username": user.username, "password": "Secure-User-Passphrase-42!"
    })
    token = login.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    assert anonymous_client.get("/api/v1/auth/me", headers=headers).status_code == 200
    logout = anonymous_client.post("/api/v1/auth/logout", headers=headers)
    assert logout.status_code == 204
    assert anonymous_client.get("/api/v1/auth/me", headers=headers).status_code == 401


def test_user_deactivation_and_role_changes_take_effect_immediately(anonymous_client, db_session):
    user = _user(db_session, "mutable-agent@example.com", role="agent")
    login = anonymous_client.post("/api/v1/auth/login", json={
        "username": user.username, "password": "Secure-User-Passphrase-42!"
    })
    headers = {"Authorization": f"Bearer {login.json()['access_token']}"}
    assert anonymous_client.get("/api/v1/users", headers=headers).status_code == 403

    user.role = "admin"
    db_session.commit()
    # A role change invalidates the old token instead of silently changing its privileges.
    assert anonymous_client.get("/api/v1/users", headers=headers).status_code == 401
    fresh_login = anonymous_client.post("/api/v1/auth/login", json={
        "username": user.username, "password": "Secure-User-Passphrase-42!"
    })
    fresh_headers = {"Authorization": f"Bearer {fresh_login.json()['access_token']}"}
    assert anonymous_client.get("/api/v1/users", headers=fresh_headers).status_code == 200

    user.is_active = False
    db_session.commit()
    assert anonymous_client.get("/api/v1/users", headers=headers).status_code == 401


def test_password_verifier_rejects_unbounded_or_malformed_scrypt_parameters():
    from app.core.security import verify_password

    assert not verify_password("anything", "scrypt$999999999$8$1$c2FsdA$ZGlnZXN0")
    assert not verify_password("anything", "scrypt$16384$8$1$c2FsdA$ZGlnZXN0")
    assert not verify_password("anything", "not-a-hash")

def test_token_without_version_claim_is_rejected(anonymous_client, db_session):
    settings = get_settings()
    legacy_token = jwt.encode({
        "sub": "1", "iat": datetime.now(timezone.utc),
        "exp": datetime.now(timezone.utc) + timedelta(minutes=5),
    }, settings.jwt_secret_key, algorithm="HS256")
    response = anonymous_client.get("/api/v1/listings", headers={"Authorization": f"Bearer {legacy_token}"})
    assert response.status_code == 401


def test_token_algorithm_allowlist_rejects_none_algorithm(anonymous_client, db_session):
    settings = get_settings()
    unsigned = jwt.encode({
        "sub": "1", "iat": datetime.now(timezone.utc),
        "exp": datetime.now(timezone.utc) + timedelta(minutes=5), "ver": 0,
    }, key="", algorithm="none")
    response = anonymous_client.get("/api/v1/listings", headers={"Authorization": f"Bearer {unsigned}"})
    assert response.status_code == 401


def test_employee_registration_requires_office_code_and_creates_agent(anonymous_client, db_session, monkeypatch):
    from app.core.config import get_settings
    monkeypatch.setattr(get_settings(), "office_registration_code", "office-secret-for-tests")
    payload = {
        "username": "new-employee",
        "password": "New-Employee-Passphrase-42!",
        "full_name": "New Employee",
        "registration_code": "wrong-code",
    }
    denied = anonymous_client.post("/api/v1/auth/register", json=payload)
    assert denied.status_code == 403
    payload["registration_code"] = "office-secret-for-tests"
    created = anonymous_client.post("/api/v1/auth/register", json=payload)
    assert created.status_code == 201, created.text
    assert created.json()["role"] == "agent"
    assert "password_hash" not in created.json()
    login = anonymous_client.post("/api/v1/auth/login", json={
        "username": payload["username"], "password": payload["password"]
    })
    assert login.status_code == 200


def test_employee_registration_cannot_assign_admin_role(anonymous_client, monkeypatch):
    from app.core.config import get_settings
    monkeypatch.setattr(get_settings(), "office_registration_code", "office-secret-for-tests")
    response = anonymous_client.post("/api/v1/auth/register", json={
        "username": "self-admin", "password": "Self-Admin-Passphrase-42!",
        "full_name": "Self Admin", "registration_code": "office-secret-for-tests", "role": "admin"
    })
    assert response.status_code == 422


def test_excel_export_requires_authentication(anonymous_client):
    response = anonymous_client.get("/api/v1/exports/listings.xlsx")
    assert response.status_code == 401
    assert response.json()["detail"] == "Giriş yapmanız gerekiyor."


def test_authenticated_excel_export_returns_valid_xlsx(client):
    from io import BytesIO
    from openpyxl import load_workbook

    response = client.get("/api/v1/exports/listings.xlsx", headers={"Authorization": "Bearer test-token-not-used-by-overridden-dependency"})
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    assert workbook.sheetnames == ["İlanlar"]
    assert workbook["İlanlar"].cell(row=1, column=1).value == "İlan ID"
    workbook.close()
